import requests
from docx import Document

# Create a simple test docx
doc = Document()
doc.add_heading('پروپوزال تست', 0)
doc.add_paragraph('این یک پروپوزال تست است برای بررسی عملکرد سیستم ارزیابی هوشمند پروپوزال‌های فارسی.')
doc.save('test_proposal.docx')

# Send to API
with open('test_proposal.docx', 'rb') as f:
    response = requests.post(
        'http://localhost:8000/api/evaluate',
        files={'file': ('test.docx', f, 'application/vnd.openxmlformats-officedocument.wordprocessingml.document')},
        params={'use_llm': False},
        timeout=30
    )

print(f'Status: {response.status_code}')
if response.status_code == 200:
    r = response.json()
    fe = r['final_evaluation']
    print(f"Score: {fe['final_score']}")
    print(f"Grade: {fe['grade']}")
    print("SUCCESS!")
else:
    print(response.text[:1000])
